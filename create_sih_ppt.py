import sys, os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def build_presentation():
    input_path = "ppt.pptx"
    output_path = "SIH2026_SIH26062_DevReact.pptx"
    
    prs = pptx.Presentation(input_path)
    print(f"Original slide count: {len(prs.slides)}")
    
    # ----------------- SLIDE 1: Title Page -----------------
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "SMART INDIA HACKATHON" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "SMART INDIA HACKATHON 2026"
                p.font.size = Pt(28)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "TITLE PAGE" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "POLARIS: Integrated Polar Expedition Logistics & Operations Suite"
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 102, 153)
            elif "Problem Statement ID" in tf.text:
                tf.clear()
                details = [
                    ("Problem Statement ID:", " SIH26062"),
                    ("Problem Statement Title:", " Integrated Polar Expedition Logistics and Asset Management System"),
                    ("Theme:", " Smart Automation | Category: Software"),
                    ("Sponsoring Organization:", " Ministry of Earth Sciences (MoES) / NCPOR, Goa"),
                    ("Team Name:", " Dev React"),
                    ("Team ID:", " 137578"),
                    ("Team Leader:", " Sparsh Sharma (sparshsharmadev@gmail.com | 9448051825)"),
                    ("Institute:", " Rajkiya Engineering College Mirzapur"),
                ]
                for idx, (label, val) in enumerate(details):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    run1 = p.add_run()
                    run1.text = label
                    run1.font.bold = True
                    run1.font.size = Pt(13)
                    run1.font.color.rgb = RGBColor(0, 51, 102)
                    
                    run2 = p.add_run()
                    run2.text = val
                    run2.font.bold = False
                    run2.font.size = Pt(13)
                    run2.font.color.rgb = RGBColor(30, 30, 30)

    # Helper function to update standard slide headers and badge
    def update_slide_base(slide, team_name="Dev React"):
        for shape in slide.shapes:
            if shape.has_text_frame:
                if "Your Team Name" in shape.text_frame.text:
                    shape.text_frame.clear()
                    p = shape.text_frame.paragraphs[0]
                    p.text = team_name
                    p.font.size = Pt(11)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(255, 255, 255)
                    p.alignment = PP_ALIGN.CENTER
                elif "@SIH Idea submission" in shape.text_frame.text:
                    shape.text_frame.clear()
                    p = shape.text_frame.paragraphs[0]
                    p.text = "SIH 2026 Idea Submission | PS: SIH26062 | MoES / NCPOR"
                    p.font.size = Pt(9)
                    p.font.color.rgb = RGBColor(120, 120, 120)

    # ----------------- SLIDE 2: Idea Title & Proposed Solution -----------------
    slide2 = prs.slides[1]
    update_slide_base(slide2)
    for shape in slide2.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "IDEA TITLE" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "POLARIS: Centralized Polar Expedition ERP & Operations Platform"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "Proposed Solution" in tf.text:
                tf.clear()
                bullets = [
                    ("Mission Context: ", "India's Antarctic (Bharati, Maitri) and Arctic (Himadri) expeditions face harsh sub-zero conditions (-40°C), severe satellite communication latencies, and high logistics costs across chartered icebreakers, polar flights, and remote field camps."),
                    ("Centralized Operational Command: ", "POLARIS provides an integrated, offline-first digital hub unifying voyage planning, cold-chain cargo tracking, station asset management, personnel safety rosters, and emergency blizzard response into a unified pane of glass."),
                    ("Core Pillars of POLARIS:", ""),
                    ("  • Expedition Voyage Planner: ", "End-to-end voyage scheduling, chartered vessel manifests (Goa ➔ Cape Town ➔ Antarctica), cargo payload balancing, and mission milestone Gantt tracking."),
                    ("  • Cold-Chain & Station Inventory: ", "Real-time tracking of mission-critical fuel (Jet A-1, Arctic diesel), life-support food rations, medical supplies, and snow vehicle (PistenBully) spare parts with automated burn-rate alerts."),
                    ("  • Personnel Roster & Movement: ", "Digital check-ins, field expedition geofencing (Schirmacher Oasis, Larsemann Hills), medical fitness clearance, and daily scientist accountability."),
                    ("  • Emergency SOS & Blizzard Cockpit: ", "1-click satellite distress dispatch, automated blizzard threat leveling (Condition 1/2/3), and dynamic Search & Rescue (SAR) / Medevac route generation."),
                    ("The 'Offline-First' Innovation: ", "Polar research stations experience frequent satellite dropouts. POLARIS runs 100% offline via local IndexedDB PWA caching and micro-syncs conflict-free when satellite/Starlink links are restored.")
                ]
                for idx, (label, val) in enumerate(bullets):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_after = Pt(4)
                    r1 = p.add_run()
                    r1.text = label
                    r1.font.bold = True
                    r1.font.size = Pt(11)
                    r1.font.color.rgb = RGBColor(0, 51, 102) if not label.startswith("  •") else RGBColor(0, 102, 153)
                    if val:
                        r2 = p.add_run()
                        r2.text = val
                        r2.font.bold = False
                        r2.font.size = Pt(11)
                        r2.font.color.rgb = RGBColor(40, 40, 40)

    # ----------------- SLIDE 3: Technical Approach -----------------
    slide3 = prs.slides[2]
    update_slide_base(slide3)
    for shape in slide3.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "TECHNICAL APPROACH" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "TECHNICAL APPROACH & SYSTEM ARCHITECTURE"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "Technologies to be used" in tf.text:
                tf.clear()
                bullets = [
                    ("Architecture Blueprint: ", "Offline-First Distributed Hub-and-Spoke Model connecting NCPOR HQ (Goa Cloud Hub) with Local Polar Edge Servers deployed at Maitri, Bharati, and chartered icebreakers."),
                    ("Technology Stack & Frameworks:", ""),
                    ("  • Frontend: ", "React 18, Vite, Vanilla CSS Design System, Responsive Glassmorphic Dashboard, Lucide Icons, Leaflet.js with Polar Stereographic Coordinate Projection."),
                    ("  • Backend & Microservices: ", "FastAPI (Python) / Node.js lightweight runtime, SQLite Edge DB + PostgreSQL Cloud DB, CRDT (Conflict-Free Replicated Data Types) for seamless multi-station sync."),
                    ("  • Local Storage & Offline Engine: ", "Progressive Web App (PWA) with Service Workers & IndexedDB for 100% local persistence during satellite blackouts."),
                    ("Automated & Intelligent Modules:", ""),
                    ("  • Predictive Fuel & Rations Burn-Rate: ", "Linear regression and seasonal trend modeling predicting resource depletion based on ambient temperature and active wintering personnel."),
                    ("  • Cargo QR/Barcode Manifest Scanning: ", "Instant container check-in/check-out for cargo offloading on fast ice."),
                    ("  • Blizzard Threat Indexing: ", "Automated weather data parsing classifying blizzard severity and locking field departure approvals."),
                    ("Security & Compliance: ", "Air-gapped encryption (AES-256), Role-Based Access Control (Expedition Leader, Medical Officer, Logistics In-Charge, Scientist), and full audit logging compliant with Madrid Environmental Protocol.")
                ]
                for idx, (label, val) in enumerate(bullets):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_after = Pt(4)
                    r1 = p.add_run()
                    r1.text = label
                    r1.font.bold = True
                    r1.font.size = Pt(11)
                    r1.font.color.rgb = RGBColor(0, 51, 102) if not label.startswith("  •") else RGBColor(0, 102, 153)
                    if val:
                        r2 = p.add_run()
                        r2.text = val
                        r2.font.bold = False
                        r2.font.size = Pt(11)
                        r2.font.color.rgb = RGBColor(40, 40, 40)

    # ----------------- SLIDE 4: Feasibility and Viability -----------------
    slide4 = prs.slides[3]
    update_slide_base(slide4)
    for shape in slide4.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "FEASIBILITY AND VIABILITY" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "FEASIBILITY, RISK MITIGATION & OPERATIONAL VIABILITY"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "Analysis of the feasibility" in tf.text:
                tf.clear()
                bullets = [
                    ("Technical Feasibility: ", "Built entirely on lightweight, resilient open web standards (PWA, IndexedDB, Docker containers). Requires zero specialized hardware — operates on rugged field laptops, tablets, and local station mini-servers with minimal CPU/RAM overhead."),
                    ("Operational Viability for NCPOR: ", "Replaces vulnerable paper logs and disconnected Excel spreadsheets with an automated, single source of truth that drastically shortens vessel turnaround times at Cape Town and Antarctic ice edges."),
                    ("Key Polar Challenges & Technical Mitigation Strategies:", ""),
                    ("  • Challenge 1 (Severe Satellite Latency & Blackouts): ", "Mitigated via Store-and-Forward delta synchronization. Transactions are committed locally in micro-seconds and synced in small compressed bursts when satellite connectivity flickers on."),
                    ("  • Challenge 2 (Extreme Environmental Hardware Risk): ", "Mitigated through dual-node local active-passive server clustering at Bharati and Maitri stations with automated hourly local flash backups."),
                    ("  • Challenge 3 (High-Stress User Ergonomics in Heavy Gear): ", "Mitigated by high-contrast polar dark-mode UI, large glove-friendly touch targets, automated barcode scanning, and 1-click critical actions."),
                    ("Scalability & Extensibility: ", "Engineered modularly so the same platform manages Antarctic Expeditions (ISEA), Arctic expeditions at Ny-Ålesund (Himadri), and high-altitude Himalayan research bases (Himansh).")
                ]
                for idx, (label, val) in enumerate(bullets):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_after = Pt(5)
                    r1 = p.add_run()
                    r1.text = label
                    r1.font.bold = True
                    r1.font.size = Pt(11)
                    r1.font.color.rgb = RGBColor(0, 51, 102) if not label.startswith("  •") else RGBColor(0, 102, 153)
                    if val:
                        r2 = p.add_run()
                        r2.text = val
                        r2.font.bold = False
                        r2.font.size = Pt(11)
                        r2.font.color.rgb = RGBColor(40, 40, 40)

    # ----------------- SLIDE 5: Impact and Benefits -----------------
    slide5 = prs.slides[4]
    update_slide_base(slide5)
    for shape in slide5.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "IMPACT AND BENEFITS" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "STRATEGIC, ECONOMIC & HUMANITARIAN IMPACT"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "Potential impact" in tf.text:
                tf.clear()
                bullets = [
                    ("Strategic National Value (Atmanirbhar Polar Governance): ", "Equips India's Ministry of Earth Sciences and NCPOR with a sovereign, indigenously developed digital backbone, reinforcing India's leadership standing within the Antarctic Treaty Consultative Meetings (ATCM)."),
                    ("Quantifiable Economic & Operational Benefits:", ""),
                    ("  • Elimination of Cargo Waste & Emergency Airfreight: ", "Prevents multi-crore losses resulting from misplaced critical components or spoiled supplies, eliminating emergency air-charter missions from Cape Town."),
                    ("  • Optimized Fuel Utilization: ", "Predictive consumption monitoring safeguards against wintering generator fuel exhaustion, ensuring unbroken life support for 14-month overwintering teams."),
                    ("  • Accelerated Icebreaker Turnaround: ", "Cuts vessel loading/unloading reconciliation time at the ice shelf from 72 hours to under 6 hours via real-time digital manifests."),
                    ("Human Safety & Life Preservation (Zero Casualties): ", "Real-time field personnel tracking, automated blizzard lockouts, and instant distress coordinate beacons slash Search & Rescue (SAR) dispatch latency from hours to under 2 minutes during deadly whiteouts."),
                    ("Environmental Compliance (Madrid Protocol): ", "Automates mandatory hazardous material tracking, fuel bunkering logs, and waste return manifests required by global Antarctic Environmental Protection accords.")
                ]
                for idx, (label, val) in enumerate(bullets):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_after = Pt(5)
                    r1 = p.add_run()
                    r1.text = label
                    r1.font.bold = True
                    r1.font.size = Pt(11)
                    r1.font.color.rgb = RGBColor(0, 51, 102) if not label.startswith("  •") else RGBColor(0, 102, 153)
                    if val:
                        r2 = p.add_run()
                        r2.text = val
                        r2.font.bold = False
                        r2.font.size = Pt(11)
                        r2.font.color.rgb = RGBColor(40, 40, 40)

    # ----------------- SLIDE 6: Research and References -----------------
    slide6 = prs.slides[5]
    update_slide_base(slide6)
    for shape in slide6.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            if "RESEARCH" in tf.text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "RESEARCH, STANDARDS & PROTOTYPE DELIVERABLES"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)
            elif "Details / Links" in tf.text:
                tf.clear()
                bullets = [
                    ("Academic & Institutional Research Foundations:", ""),
                    ("  1. National Centre for Polar and Ocean Research (NCPOR): ", "Indian Scientific Expedition to Antarctica (ISEA) Logistics & Station Operations Documentation (41st - 44th Expeditions)."),
                    ("  2. COMNAP Guidelines: ", "Council of Managers of National Antarctic Programs — Standards for Antarctic Operations, Field Safety & Waste Management Protocols."),
                    ("  3. SCAR (Scientific Committee on Antarctic Research): ", "Best Practices in Environmental Data Collection & Multilateral Polar Logistics Coordination."),
                    ("  4. The Protocol on Environmental Protection to the Antarctic Treaty (Madrid Protocol): ", "Hazardous Waste Classification & Fuel Deposition Compliance Guidelines."),
                    ("Working Prototype & Verification Deliverables:", ""),
                    ("  • Live Interactive Web Application: ", "Fully functional React + Vite Operations Cockpit featuring Voyage Gantt, Station Gauges, Cargo Inventory, and SOS Dispatch simulator."),
                    ("  • Prototype Demo & Video Walkthrough: ", "Video showcasing the end-to-end voyage creation, offline barcode scan, and blizzard response drill."),
                    ("  • Open-Source Code Repository: ", "Complete codebase, PWA offline service worker configuration, and mock polar dataset on Team Dev React GitHub."),
                    ("Team Commitment: ", "Team Dev React is prepared with a functional deployment, architecture documentation, and live simulation for the SIH Grand Finale.")
                ]
                for idx, (label, val) in enumerate(bullets):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_after = Pt(4)
                    r1 = p.add_run()
                    r1.text = label
                    r1.font.bold = True
                    r1.font.size = Pt(11)
                    r1.font.color.rgb = RGBColor(0, 51, 102) if not label.startswith("  ") else RGBColor(0, 102, 153)
                    if val:
                        r2 = p.add_run()
                        r2.text = val
                        r2.font.bold = False
                        r2.font.size = Pt(11)
                        r2.font.color.rgb = RGBColor(40, 40, 40)

    # ----------------- Remove Slide 7 (Instructions slide) -----------------
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Removed instruction slide 7. Total slides now: 6")

    prs.save(output_path)
    # Also overwrite ppt.pptx so the user has it in their primary file
    prs.save("ppt.pptx")
    print(f"Successfully saved generated presentation to {output_path} and ppt.pptx!")

if __name__ == "__main__":
    build_presentation()
