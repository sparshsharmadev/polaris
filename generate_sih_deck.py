from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize Presentation with 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
BG_DARK = RGBColor(11, 19, 43)        # #0B132B Polar Deep Navy
CARD_BG = RGBColor(28, 37, 65)        # #1C2541 Card Slate Navy
ACCENT_CYAN = RGBColor(91, 192, 190)   # #5BC0BE Cyan
ACCENT_BLUE = RGBColor(77, 144, 254)   # #4D90FE Ice Blue
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(180, 195, 215)
BORDER_COLOR = RGBColor(58, 80, 107)

def apply_slide_background(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_DARK
    bg_shape.line.fill.background()

def add_header(slide, category_text, title_text):
    # Category / Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_text.upper()
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

def add_footer(slide):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026 | Team ID: 137578 | PS ID: SIH26062 | MoES / NCPOR"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, title, content_bullets):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = BORDER_COLOR
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.space_after = Pt(10)

    for bullet in content_bullets:
        p = tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(6)

# ==========================================
# SLIDE 1: COVER SLIDE
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide1)

# Badge
badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(3.2), Inches(0.4))
badge.fill.solid()
badge.fill.fore_color.rgb = CARD_BG
badge.line.color.rgb = ACCENT_CYAN
tf_b = badge.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "SMART INDIA HACKATHON 2026"
p_b.alignment = PP_ALIGN.CENTER
p_b.font.size = Pt(11)
p_b.font.bold = True
p_b.font.color.rgb = ACCENT_CYAN

# Title
t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.8))
tf = t_box.text_frame
tf.word_wrap = True
p1 = tf.paragraphs[0]
p1.text = "POLARIS"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE

p2 = tf.add_paragraph()
p2.text = "Integrated Polar Expedition Logistics & Operations Suite"
p2.font.size = Pt(20)
p2.font.color.rgb = ACCENT_CYAN
p2.space_before = Pt(6)

# Metadata Card Left
add_card(slide1, Inches(0.8), Inches(3.4), Inches(5.6), Inches(3.2), "Problem Statement Scope", [
    "Problem Statement ID: SIH26062",
    "Title: Integrated Polar Expedition Logistics and Asset Management System",
    "Category: Software | Theme: Smart Automation",
    "Sponsoring Agency: Ministry of Earth Sciences (MoES) / NCPOR, Goa"
])

# Metadata Card Right
add_card(slide1, Inches(6.8), Inches(3.4), Inches(5.7), Inches(3.2), "Team Credentials", [
    "Team Name: Dev React (Team ID: 137578)",
    "Institute: Rajkiya Engineering College Mirzapur",
    "Team Leader: Sparsh Sharma",
    "Contact: sparshsharmadev@gmail.com | +91 9448051825"
])
add_footer(slide1)

# ==========================================
# SLIDE 2: PROPOSED SOLUTION & PILLARS
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide2)
add_header(slide2, "Concept Overview", "POLARIS: Centralized Polar Expedition ERP & Operations Platform")

# Top Context Banner
ctx_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0))
ctx_card.fill.solid()
ctx_card.fill.fore_color.rgb = CARD_BG
ctx_card.line.color.rgb = ACCENT_BLUE
tf_c = ctx_card.text_frame
tf_c.word_wrap = True
tf_c.margin_left = Inches(0.2)
tf_c.margin_top = Inches(0.12)
p_c = tf_c.paragraphs[0]
p_c.text = "Operational Reality:"
p_c.font.bold = True
p_c.font.size = Pt(11)
p_c.font.color.rgb = ACCENT_CYAN
p_c2 = tf_c.add_paragraph()
p_c2.text = "India's Antarctic (Bharati, Maitri) and Arctic (Himadri) expeditions face sub-zero conditions (-40°C), severe satellite blackouts, and disconnected spreadsheets across 14-month overwintering cycles. POLARIS delivers a resilient, offline-first operations hub."
p_c2.font.size = Pt(10)
p_c2.font.color.rgb = TEXT_WHITE

# 4 Pillars
add_card(slide2, Inches(0.8), Inches(2.7), Inches(2.7), Inches(4.0), "Voyage Planning", [
    "Chartered vessel tracking (Goa → Cape Town → Ice Edge).",
    "Cargo payload balancing.",
    "Mission Gantt milestones.",
    "Demurrage minimization."
])

add_card(slide2, Inches(3.8), Inches(2.7), Inches(2.7), Inches(4.0), "Cold-Chain Asset ERP", [
    "Arctic diesel & Jet A-1 telemetry.",
    "Perishable ration expiry pacing.",
    "PistenBully vehicle parts logs.",
    "Automated burn-rate alerts."
])

add_card(slide2, Inches(6.8), Inches(2.7), Inches(2.7), Inches(4.0), "Personnel & Safety", [
    "Geofenced traverse check-ins (Schirmacher / Larsemann).",
    "Overdue departure lockouts.",
    "Daily digital muster roll.",
    "Scientist fitness clearances."
])

add_card(slide2, Inches(9.8), Inches(2.7), Inches(2.7), Inches(4.0), "Offline SOS & Blizzard", [
    "100% PWA offline IndexedDB.",
    "Condition 1/2/3 weather indexing.",
    "1-click distress satellite packet.",
    "Automated SAR routing."
])
add_footer(slide2)

# ==========================================
# SLIDE 3: TECHNICAL APPROACH & TOPOLOGY
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide3)
add_header(slide3, "System Engineering", "Distributed Hub-and-Spoke Topology & Tech Architecture")

# Left Column: Topology Architecture Box
arch_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
arch_box.fill.solid()
arch_box.fill.fore_color.rgb = CARD_BG
arch_box.line.color.rgb = ACCENT_CYAN
tf_a = arch_box.text_frame
tf_a.word_wrap = True
tf_a.margin_left = Inches(0.25)
tf_a.margin_top = Inches(0.25)
p_at = tf_a.paragraphs[0]
p_at.text = "OFFLINE-FIRST SYNCHRONIZATION TOPOLOGY"
p_at.font.bold = True
p_at.font.size = Pt(13)
p_at.font.color.rgb = ACCENT_CYAN
p_at.space_after = Pt(12)

steps = [
    ("[ FIELD TRAVERSES & SHORE OUTPOSTS ]\nRugged Tablets / Scanners / IndexedDB Local Storage"),
    ("                     │ (Local PWA commit in < 5ms)"),
    ("                     ▼"),
    ("[ STATION / SHIP EDGE GATEWAY ]\nLocal Docker Node (Bharati / Maitri / Icebreaker)\nSQLite + CRDTs (Conflict-Free Replicated Data)"),
    ("                     │ (Intermittent Satellite / Starlink Uplink)"),
    ("                     ▼ (Delta Binary Packets / Compressed Sync)"),
    ("[ NCPOR HQ CENTRAL CLOUD (GOA) ]\nPostgreSQL Master Database + Madrid Protocol Audit Vault")
]
for step in steps:
    p = tf_a.add_paragraph()
    p.text = step
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_WHITE
    p.font.name = "Courier New"
    p.space_after = Pt(4)

# Right Cards
add_card(slide3, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.45), "Core Technology Stack", [
    "Frontend: React 18, Vite, Tailwind CSS, Lucide Icons.",
    "Polar GIS: Leaflet.js with EPSG:3031 Polar Stereographic Projection.",
    "Edge Runtime: Python FastAPI + SQLite lightweight local instances.",
    "Cloud Hub: PostgreSQL, Docker containers, Redis queue worker."
])

add_card(slide3, Inches(6.8), Inches(4.25), Inches(5.7), Inches(2.45), "Key Engineering Features", [
    "CRDT Conflict Resolution: Automatically merges multi-device offline updates.",
    "Predictive Burn Pacing: Dynamic consumption models tied to weather severity.",
    "Air-Gapped Encryption: AES-256 local volume protection & RBAC.",
    "Madrid Protocol Compliance: Audited records for fuel and hazardous waste."
])
add_footer(slide3)

# ==========================================
# SLIDE 4: FEASIBILITY & MITIGATION TABLE
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide4)
add_header(slide4, "Operational Viability", "Feasibility Analysis & Polar Risk Mitigation Strategies")

# Table Setup
rows, cols = 5, 3
left = Inches(0.8)
top = Inches(1.6)
width = Inches(11.7)
height = Inches(4.9)

table_shape = slide4.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table
table.columns[0].width = Inches(3.0)
table.columns[1].width = Inches(3.7)
table.columns[2].width = Inches(5.0)

headers = ["Polar Environmental Challenge", "Operational Vulnerability", "POLARIS Architectural Mitigation"]
for idx, header in enumerate(headers):
    cell = table.cell(0, idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BG
    cell.text = header
    p = cell.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = ACCENT_CYAN

data = [
    ("Satellite Outages & Latency", "Cloud-only dashboards fail, losing station logs.", "Store-and-Forward PWA architecture: mutations persist in local IndexedDB in <5ms and batch sync over delta bursts."),
    ("Hardware Failure in Extreme Cold", "Server drive freeze corrupts mission manifest records.", "Active-Passive dual clustering on rugged edge laptops with automated flash storage backups every hour."),
    ("Sub-Zero Ergonomics (-40°C)", "Small buttons unusable while wearing bulky mittens.", "Polar Dark UI: 48px+ touch targets, high-contrast displays, and rapid hardware barcode/QR scanning."),
    ("Strict Antarctic Treaty Accords", "Manual logs risk hazardous waste audit fines.", "Automated Madrid Protocol reporting tracks hazardous materials and cargo returns from ice shelves.")
]

for row_idx, (col0, col1, col2) in enumerate(data, start=1):
    for col_idx, text in enumerate([col0, col1, col2]):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG_DARK
        cell.text = text
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_WHITE

add_footer(slide4)

# ==========================================
# SLIDE 5: STRATEGIC & ECONOMIC IMPACT
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide5)
add_header(slide5, "Value Proposition", "Strategic, Economic & Humanitarian Impact for MoES / NCPOR")

# 4 Metric Cards
add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4), "Vessel Turnaround: 72h → <6h", [
    "Accelerates ice-edge unloading manifest verification.",
    "Saves multi-crore chartered icebreaker demurrage.",
    "Real-time scanning reconciles shipping pallets on fast ice."
])

add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.4), "Elimination of Emergency Air-Charters", [
    "Accurate seasonal fuel & parts burn-rate forecasting.",
    "Prevents emergency Cape Town flight missions.",
    "Protects uninterrupted 14-month overwintering life support."
])

add_card(slide5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.4), "Human Safety: Zero Casualties", [
    "Automatic Condition 1/2/3 blizzard lockout triggers.",
    "Distress coordinate packets dispatch in <2 minutes.",
    "Geofenced traverse tracking prevents whiteout disorientation."
])

add_card(slide5, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.4), "Sovereign Polar Governance", [
    "Indigenous tech backbone for polar stations (Bharati/Maitri).",
    "Reinforces Indian leadership at Antarctic Treaty meetings.",
    "Standardized across Antarctic, Arctic, and Himalayan bases."
])
add_footer(slide5)

# ==========================================
# SLIDE 6: STANDARDS & FINALE DELIVERABLES
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
apply_slide_background(slide6)
add_header(slide6, "Execution Readiness", "Operational Standards & Grand Finale Prototype Deliverables")

add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), "Institutional Research Standards", [
    "NCPOR ISEA Documentation: Operational frameworks from 39th to 44th Indian Scientific Expeditions.",
    "COMNAP Guidelines: Council of Managers of National Antarctic Programs safety & logistics protocols.",
    "SCAR Framework: Best practices for multilateral polar environmental data collection.",
    "Madrid Protocol (Annex III & IV): Mandatory compliance tracking for waste disposal and marine fuel bunkering."
])

add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Evaluation Deliverables Ready for Demo", [
    "Live Interactive Web Cockpit: Fully functional React + Vite app with Gantt schedules, tank gauges, and SOS panels.",
    "Zero-Bandwidth Blackout Simulation: Live demonstration showing network cut, local transactions, and conflict-free CRDT sync.",
    "Clean Open-Source Repository: Production-ready Docker edge setup, PWA service workers, and polar seed datasets on GitHub.",
    "Ready for National Grand Finale: Fully configured for 36-hour physical hackathon deployment."
])
add_footer(slide6)

# Save
output_path = "SIH2026_SIH26062_DevReact.pptx"
prs.save(output_path)
print(f"Presentation saved successfully to: {output_path}")